"""Excel export functionality for Identity CLI using openpyxl."""

from datetime import datetime
from pathlib import Path
from typing import List

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from identity_cli.models import Identity


def get_unique_export_path(exports_dir: Path) -> Path:
    """Generate a unique timestamped filename for the Excel export.

    Never silently overwrites an existing file.
    """
    timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    base_name = f"identity_{timestamp}"
    candidate = exports_dir / f"{base_name}.xlsx"

    counter = 1
    while candidate.exists():
        candidate = exports_dir / f"{base_name}_{counter}.xlsx"
        counter += 1

    return candidate


def export_to_excel(identities: List[Identity], exports_dir: Path) -> Path:
    """Export a list of identities to an Excel (.xlsx) file.

    Includes clean formatting:
    - Headers on row 1 (frozen)
    - Auto-adjusted column widths
    - Optional phone column based on identity contents
    """
    exports_dir.mkdir(parents=True, exist_ok=True)
    file_path = get_unique_export_path(exports_dir)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Identities"

    # Determine if any identity has a phone number
    has_phone = any(identity.phone is not None for identity in identities)

    headers = ["ID", "Name", "Email", "Password"]
    if has_phone:
        headers.append("Phone")

    ws.append(headers)

    # Make header bold and freeze top row
    header_font = Font(bold=True)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font

    ws.freeze_panes = "A2"

    # Append rows
    for idx, identity in enumerate(identities, start=1):
        row = [idx, identity.name, identity.email, identity.password]
        if has_phone:
            row.append(identity.phone or "")
        ws.append(row)

    # Adjust column widths automatically
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(file_path)
    return file_path
