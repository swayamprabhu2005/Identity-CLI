"""Export functionality for Identity CLI (Excel, JSON, CSV)."""

import csv
import json
from pathlib import Path
from typing import List, Optional

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from identity_cli.models import Identity

FIELD_LABELS = {
    "name": "Name",
    "email": "Email",
    "password": "Password",
    "phone": "Phone",
}


def get_safe_export_path(target_dir: Path, base_name: str = "identities", ext: str = "xlsx") -> Path:
    """Generate a collision-safe path: identities.ext, identities_1.ext, identities_2.ext, etc."""
    target_dir.mkdir(parents=True, exist_ok=True)
    candidate = target_dir / f"{base_name}.{ext}"
    if not candidate.exists():
        return candidate
    counter = 1
    while candidate.exists():
        candidate = target_dir / f"{base_name}_{counter}.{ext}"
        counter += 1
    return candidate


def export_to_excel(
    identities: List[Identity],
    exports_dir: Path,
    fields: Optional[List[str]] = None
) -> Path:
    """Export a list of identities to an Excel (.xlsx) file with only selected fields."""
    exports_dir.mkdir(parents=True, exist_ok=True)
    file_path = get_safe_export_path(exports_dir, base_name="identities", ext="xlsx")

    if fields is not None:
        active_fields = [f for f in fields if f in FIELD_LABELS]
    else:
        active_fields = []
        if any(i.name is not None for i in identities):
            active_fields.append("name")
        if any(i.email is not None for i in identities):
            active_fields.append("email")
        if any(i.password is not None for i in identities):
            active_fields.append("password")
        if any(i.phone is not None for i in identities):
            active_fields.append("phone")
        if not active_fields:
            active_fields = ["name", "email", "password"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Identities"

    headers = ["ID"] + [FIELD_LABELS[f] for f in active_fields]
    ws.append(headers)

    header_font = Font(bold=True)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font

    ws.freeze_panes = "A2"

    for idx, identity in enumerate(identities, start=1):
        row = [idx]
        for field in active_fields:
            row.append(getattr(identity, field, None) or "")
        ws.append(row)

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(file_path)
    return file_path


def export_to_json(
    identities: List[Identity],
    exports_dir: Path,
    fields: Optional[List[str]] = None
) -> Path:
    """Export identities to a JSON file containing only selected fields."""
    exports_dir.mkdir(parents=True, exist_ok=True)
    file_path = get_safe_export_path(exports_dir, base_name="identities", ext="json")

    data = [item.to_dict(selected_fields=fields) for item in identities]
    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)

    return file_path


def export_to_csv(
    identities: List[Identity],
    exports_dir: Path,
    fields: Optional[List[str]] = None
) -> Path:
    """Export identities to a CSV file containing only selected fields."""
    exports_dir.mkdir(parents=True, exist_ok=True)
    file_path = get_safe_export_path(exports_dir, base_name="identities", ext="csv")

    if fields is not None:
        active_fields = [f for f in fields if f in FIELD_LABELS]
    else:
        active_fields = []
        if any(i.name is not None for i in identities):
            active_fields.append("name")
        if any(i.email is not None for i in identities):
            active_fields.append("email")
        if any(i.password is not None for i in identities):
            active_fields.append("password")
        if any(i.phone is not None for i in identities):
            active_fields.append("phone")
        if not active_fields:
            active_fields = ["name", "email", "password"]

    headers = ["ID"] + [FIELD_LABELS[f] for f in active_fields]

    with open(file_path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for idx, identity in enumerate(identities, start=1):
            row = [idx]
            for field in active_fields:
                row.append(getattr(identity, field, None) or "")
            writer.writerow(row)

    return file_path
