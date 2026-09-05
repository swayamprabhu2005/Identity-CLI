"""Tests for Excel export."""

from pathlib import Path
import openpyxl
from identity_cli.excel import export_to_excel, get_unique_export_path
from identity_cli.models import Identity


def test_excel_export_without_phone(temp_dir: Path) -> None:
    exports_dir = temp_dir / "exports"
    identities = [
        Identity(
            name="Ethan Brooks",
            email="ethan.brooks482@example.com",
            password="V!7qL@92xK#p",
        ),
        Identity(
            name="Sarah Mitchell",
            email="sarah.mitchell731@example.com",
            password="mR8$zT21!qW",
        ),
    ]

    file_path = export_to_excel(identities, exports_dir)

    assert file_path.exists()
    assert file_path.suffix == ".xlsx"

    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    # Check headers
    headers = [cell.value for cell in ws[1]]
    assert headers == ["ID", "Name", "Email", "Password"]
    assert "Phone" not in headers

    # Check rows
    assert ws.max_row == 3  # Header + 2 rows
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=2, column=2).value == "Ethan Brooks"
    assert ws.cell(row=3, column=1).value == 2
    assert ws.cell(row=3, column=2).value == "Sarah Mitchell"

    # Check freeze panes
    assert ws.freeze_panes == "A2"


def test_excel_export_with_phone(temp_dir: Path) -> None:
    exports_dir = temp_dir / "exports"
    identities = [
        Identity(
            name="Maya Richardson",
            email="maya.richardson731@example.com",
            password="pQ8!xR2#Lm91",
            phone="+1 202-555-0103",
        )
    ]

    file_path = export_to_excel(identities, exports_dir)
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    assert headers == ["ID", "Name", "Email", "Password", "Phone"]
    assert ws.cell(row=2, column=5).value == "+1 202-555-0103"


def test_excel_collision_safe_naming(temp_dir: Path) -> None:
    exports_dir = temp_dir / "exports"
    exports_dir.mkdir(parents=True, exist_ok=True)

    # First call
    path1 = get_unique_export_path(exports_dir)
    path1.touch()

    # Second call for the same moment in time should produce a non-colliding name
    path2 = get_unique_export_path(exports_dir)
    assert path1 != path2
    assert not path2.exists()
